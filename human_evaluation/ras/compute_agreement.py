#!/usr/bin/env python3
"""Compute agreement for support-only per-qid RAS workbooks."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = SCRIPT_DIR / "results"
DEFAULT_SAMPLES = SCRIPT_DIR / "samples.json"
WORKBOOK_GLOB = "qid_*.xlsx"


def yes_no_to_bool(value: object) -> bool | None:
    text = " ".join(str(value or "").strip().lower().split())
    if text == "yes":
        return True
    if text == "no":
        return False
    return None


def normalize_gold(value: object) -> bool | None:
    if isinstance(value, bool):
        return value
    text = " ".join(str(value or "").strip().lower().split())
    if text in {"true", "yes", "1"}:
        return True
    if text in {"false", "no", "0"}:
        return False
    return None


def load_support_gold(samples_path: Path) -> dict[tuple[str, int], bool]:
    with samples_path.open("r", encoding="utf-8") as f:
        samples = json.load(f)

    gold = {}
    for qids in samples.values():
        for qid, item in qids.items():
            for support_idx, (_, value) in enumerate(item.get("support", {}).items(), start=1):
                gold[(qid, support_idx)] = bool(value)
    return gold


def workbook_records(path: Path, gold_lookup: dict[tuple[str, int], bool]) -> dict[tuple[str, int], dict[str, bool | None]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    records = {}

    for sheet_name in wb.sheetnames:
        prefix = f"{path.stem}_support_"
        if not sheet_name.startswith(prefix):
            continue

        ws = wb[sheet_name]
        support_idx = int(sheet_name[len(prefix):])
        key = (path.stem, support_idx)
        records[key] = {
            "gold": gold_lookup.get(key),
            "label": yes_no_to_bool(ws["B7"].value),
        }

    if not records:
        raise ValueError(f"{path.name}: no {path.stem}_support_* sheets found")
    return records


def annotator_records(
    annotator_dir: Path,
    workbook_glob: str,
    gold_lookup: dict[tuple[str, int], bool],
) -> dict[tuple[str, int], dict[str, bool | None]]:
    records = {}
    paths = sorted(annotator_dir.glob(workbook_glob))
    if not paths:
        raise FileNotFoundError(f"No qid workbooks matched {annotator_dir / workbook_glob}")

    for path in paths:
        records.update(workbook_records(path, gold_lookup))
    return records


def percentage_agreement(labels_a: dict, labels_b: dict) -> tuple[int, int, float]:
    keys = sorted(set(labels_a) & set(labels_b))
    total = 0
    matches = 0
    for key in keys:
        a = labels_a[key]
        b = labels_b[key]
        if not isinstance(a, bool) or not isinstance(b, bool):
            continue
        total += 1
        matches += int(a == b)
    return total, matches, matches / total if total else 0.0


def fleiss_kappa_binary(label_sets: list[dict], item_keys: list[tuple[str, int]]) -> tuple[float | None, int]:
    rows = []
    for key in item_keys:
        votes = [labels.get(key) for labels in label_sets]
        if not all(isinstance(vote, bool) for vote in votes):
            continue
        true_count = sum(1 for vote in votes if vote)
        false_count = len(votes) - true_count
        rows.append((true_count, false_count))

    item_count = len(rows)
    annotator_count = len(label_sets)
    if item_count == 0:
        return None, 0

    p_true = sum(row[0] for row in rows) / (item_count * annotator_count)
    p_false = sum(row[1] for row in rows) / (item_count * annotator_count)
    p_expected = p_true**2 + p_false**2
    p_observed = sum(
        (true_count * (true_count - 1) + false_count * (false_count - 1))
        / (annotator_count * (annotator_count - 1))
        for true_count, false_count in rows
    ) / item_count

    if abs(1 - p_expected) < 1e-12:
        return (1.0 if abs(p_observed - 1.0) < 1e-12 else 0.0), item_count
    return (p_observed - p_expected) / (1 - p_expected), item_count


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Compute support-only RAS agreement.")
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--samples", type=Path, default=DEFAULT_SAMPLES)
    parser.add_argument("--workbook-glob", default=WORKBOOK_GLOB)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    annotator_dirs = [path for path in sorted(args.input_dir.iterdir()) if path.is_dir()]
    if not annotator_dirs:
        raise FileNotFoundError(f"No annotator directories found in {args.input_dir}")

    gold_lookup = load_support_gold(args.samples)
    workbooks = {
        annotator_dir.name: annotator_records(annotator_dir, args.workbook_glob, gold_lookup)
        for annotator_dir in annotator_dirs
    }
    label_sets = {
        name: {key: record["label"] for key, record in records.items()}
        for name, records in workbooks.items()
    }
    gold = {
        key: record["gold"]
        for records in workbooks.values()
        for key, record in records.items()
    }

    print("annotator,total,matches_with_gold,agreement_percent,missing_labels")
    for name, labels in label_sets.items():
        total, matches, agreement = percentage_agreement(labels, gold)
        missing = sum(1 for label in labels.values() if not isinstance(label, bool))
        print(f"{name},{total},{matches},{100 * agreement:.2f},{missing}")

    print()
    print("pair,total,matches,agreement_percent")
    for left, right in combinations(label_sets, 2):
        total, matches, agreement = percentage_agreement(label_sets[left], label_sets[right])
        print(f"{left}__vs__{right},{total},{matches},{100 * agreement:.2f}")

    all_item_keys = sorted(set().union(*(set(labels) for labels in label_sets.values())))
    kappa, kappa_items = fleiss_kappa_binary(list(label_sets.values()), all_item_keys)
    kappa_text = "" if kappa is None else f"{kappa:.4f}"
    print()
    print(f"fleiss_kappa_items,{kappa_items}")
    print(f"fleiss_kappa,{kappa_text}")

    majority_labels = {}
    for key in all_item_keys:
        valid = [labels.get(key) for labels in label_sets.values() if isinstance(labels.get(key), bool)]
        counts = Counter(valid)
        if counts[True] > counts[False]:
            majority_labels[key] = True
        elif counts[False] > counts[True]:
            majority_labels[key] = False

    total, matches, agreement = percentage_agreement(majority_labels, gold)
    print(f"majority_vs_gold,{total},{matches},{100 * agreement:.2f}")


if __name__ == "__main__":
    main()
