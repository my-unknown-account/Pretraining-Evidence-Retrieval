#!/usr/bin/env python3
"""Create one support-only RAS human-evaluation workbook per qid."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_SAMPLES = SCRIPT_DIR / "samples.json"
DEFAULT_OUTPUT_DIR = SCRIPT_DIR / "workbooks"
SAMPLE_SIZE = 300

LABELS = ["Yes", "No"]


def clean_text(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def load_samples(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def build_qid_rows(samples_path: Path, limit: int) -> dict[str, dict[str, object]]:
    samples = load_samples(samples_path)
    qid_rows: dict[str, dict[str, object]] = {}
    total = 0

    for model, qids in samples.items():
        for qid, item in qids.items():
            rows = []
            for support_idx, (passage, gold) in enumerate(item.get("support", {}).items(), start=1):
                rows.append(
                    {
                        "support_idx": support_idx,
                        "subject": clean_text(item.get("sub")),
                        "relation": clean_text(item.get("rel")),
                        "object": clean_text(item.get("obj")),
                        "passage": clean_text(passage),
                        "gold": bool(gold),
                    }
                )
                total += 1
                if total == limit:
                    qid_rows[qid] = {"model": model, "rows": rows}
                    return qid_rows

            qid_rows[qid] = {"model": model, "rows": rows}

    if total < limit:
        raise ValueError(f"requested {limit} support passages, found only {total}")
    return qid_rows


def add_label_validation(ws) -> None:
    formula = '"' + ",".join(LABELS) + '"'
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Invalid Input",
        error="Please choose only Yes or No.",
    )
    ws.add_data_validation(validation)
    validation.add(ws["B7"])


def style_sheet(ws) -> None:
    title_fill = PatternFill("solid", fgColor="0F172A")
    label_fill = PatternFill("solid", fgColor="DBEAFE")
    text_fill = PatternFill("solid", fgColor="F8FAFC")
    answer_fill = PatternFill("solid", fgColor="FEF3C7")
    title_font = Font(color="FFFFFF", bold=True, size=14)
    key_font = Font(color="1E3A8A", bold=True)
    body_font = Font(size=11)
    answer_font = Font(color="92400E", bold=True)
    thin_gray = Side(style="thin", color="CBD5E1")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws["A1"].fill = title_fill
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=3, max_row=7, max_col=2):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = text_fill

    for row_idx in range(3, 8):
        ws.cell(row=row_idx, column=1).fill = label_fill
        ws.cell(row=row_idx, column=1).font = key_font

    ws["A7"].fill = answer_fill
    ws["A7"].font = answer_font
    ws["B7"].fill = answer_fill
    ws["B7"].font = answer_font
    ws["B7"].alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 22,
        "B": 200,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 30
    for idx in [3, 4, 5]:
        ws.row_dimensions[idx].height = 24
    ws.row_dimensions[6].height = 500
    ws.row_dimensions[7].height = 30


def add_passage_sheet(ws, qid: str, model: str, row: dict[str, object]) -> None:
    ws.merge_cells("A1:B1")
    ws["A1"] = "Support Annotation"

    values = [
        ("Subject", row["subject"]),
        ("Relation", row["relation"]),
        ("Object", row["object"]),
        ("Text", row["passage"]),
        ("Is it supported?", ""),
    ]
    for row_idx, (key, value) in enumerate(values, start=3):
        ws.cell(row=row_idx, column=1, value=key)
        ws.cell(row=row_idx, column=2, value=value)

    add_label_validation(ws)
    style_sheet(ws)


def create_qid_workbook(
    qid: str,
    model: str,
    rows: list[dict[str, object]],
    output_dir: Path,
) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for row in rows:
        ws = wb.create_sheet(f"{qid}_support_{int(row['support_idx'])}"[:31])
        add_passage_sheet(ws, qid, model, row)

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{qid}.xlsx"
    wb.save(path)
    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create support-only RAS qid workbooks.")
    parser.add_argument("--samples", type=Path, default=DEFAULT_SAMPLES)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--limit", type=int, default=SAMPLE_SIZE)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    qid_rows = build_qid_rows(args.samples, args.limit)
    outputs = [
        create_qid_workbook(qid, str(data["model"]), list(data["rows"]), args.output_dir)
        for qid, data in qid_rows.items()
    ]

    print(f"[ok] wrote {len(outputs)} qid workbooks with {args.limit} support passages")
    for output in outputs:
        print(f"[ok] wrote {output}")


if __name__ == "__main__":
    main()
