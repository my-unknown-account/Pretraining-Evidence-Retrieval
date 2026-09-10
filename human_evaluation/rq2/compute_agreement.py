#!/usr/bin/env python3
"""Compute agreement between filled annotation workbooks and GPT judge labels."""

from __future__ import annotations

import argparse
import csv
import random
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parents[1]
DEFAULT_INPUT_DIR = SCRIPT_DIR / "results"
DEFAULT_RQ2_EXAMPLES = DEFAULT_INPUT_DIR / "rq2_examples.csv"
DEFAULT_SEED = 20260909

VALID_LABELS = ["TRUE_OBJECT", "FALSE_CONTEXT_OBJECT", "OTHER"]
VALID_LABEL_SET = set(VALID_LABELS)
SHEETS_PER_WORKBOOK = 10
ROWS_PER_SHEET = 30
SAMPLE_SIZE = SHEETS_PER_WORKBOOK * ROWS_PER_SHEET
WORKBOOK_GLOB = "rq2_human_evaluation_annotator_*_filled.xlsx"


def normalize_label(value: object) -> str:
    text = " ".join(str(value or "").strip().split()).upper()
    if text in VALID_LABEL_SET:
        return text
    for label in VALID_LABELS:
        if label in text:
            return label
    return ""


def read_rq2_examples(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def expected_gpt_labels(rq2_examples: Path, seed: int) -> list[str]:
    rows = read_rq2_examples(rq2_examples)
    if SAMPLE_SIZE > len(rows):
        raise ValueError(f"sample size {SAMPLE_SIZE} exceeds available rows {len(rows)}")
    sampled = random.Random(seed).sample(rows, SAMPLE_SIZE)
    labels = [normalize_label(row.get("negative_label")) for row in sampled]
    invalid = [idx + 1 for idx, label in enumerate(labels) if not label]
    if invalid:
        raise ValueError(f"invalid GPT labels at sampled positions: {invalid[:10]}")
    return labels


def workbook_labels(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    labels = []
    for sheet_name in [f"Batch {idx:02d}" for idx in range(1, SHEETS_PER_WORKBOOK + 1)]:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"{path.name}: missing sheet {sheet_name}")
        ws = wb[sheet_name]
        headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
        if "label" not in headers:
            raise ValueError(f"{path.name}:{sheet_name}: missing label column")
        label_col = headers.index("label") + 1
        for row_idx in range(2, ROWS_PER_SHEET + 2):
            labels.append(normalize_label(ws.cell(row=row_idx, column=label_col).value))
    return labels


def compute_agreement(human_labels: list[str], gpt_labels: list[str]) -> dict[str, object]:
    if len(human_labels) != len(gpt_labels):
        raise ValueError(f"label count mismatch: {len(human_labels)} vs {len(gpt_labels)}")

    missing = [idx + 1 for idx, label in enumerate(human_labels) if label not in VALID_LABEL_SET]
    if missing:
        raise ValueError(f"missing/invalid workbook labels at positions: {missing[:20]}")

    total = len(gpt_labels)
    matches = sum(human == gpt for human, gpt in zip(human_labels, gpt_labels))
    disagreements = Counter(
        (human, gpt)
        for human, gpt in zip(human_labels, gpt_labels)
        if human != gpt
    )
    return {
        "total": total,
        "matches": matches,
        "agreement": matches / total,
        "disagreements": disagreements,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compute percentage agreement for RQ2 filled annotation workbooks."
    )
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--rq2-examples", type=Path, default=DEFAULT_RQ2_EXAMPLES)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--workbook-glob", default=WORKBOOK_GLOB)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    gpt_labels = expected_gpt_labels(args.rq2_examples, args.seed)
    paths = sorted(args.input_dir.glob(args.workbook_glob))
    if not paths:
        raise FileNotFoundError(f"No filled workbooks matched {args.input_dir / args.workbook_glob}")

    overall_matches = 0
    overall_total = 0
    print("workbook,total,matches,agreement_percent")
    for path in paths:
        labels = workbook_labels(path)
        result = compute_agreement(labels, gpt_labels)
        overall_matches += int(result["matches"])
        overall_total += int(result["total"])
        print(
            f"{path.name},{result['total']},{result['matches']},"
            f"{100 * float(result['agreement']):.2f}"
        )

    overall = overall_matches / overall_total if overall_total else 0
    print(f"OVERALL,{overall_total},{overall_matches},{100 * overall:.2f}")


if __name__ == "__main__":
    main()
