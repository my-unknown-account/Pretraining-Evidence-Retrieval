#!/usr/bin/env python3
"""Create Excel workbooks for RQ2 human annotation."""

from __future__ import annotations

import argparse
import csv
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RQ2_EXAMPLES = REPO_ROOT / "scripts" / "analysis" / "outputs" / "rq2" / "rq2_examples.csv"
DEFAULT_DATASET = REPO_ROOT / "data" / "dataset.json"
DEFAULT_CONTRADICTORY_DIR = REPO_ROOT / "scripts" / "inference" / "contradictory_context" / "results"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "human_evaluation" / "rq2" / "results"
DEFAULT_SEED = 20260909

LABELS = ["TRUE_OBJECT", "FALSE_CONTEXT_OBJECT", "OTHER"]
SHEETS_PER_WORKBOOK = 10
ROWS_PER_SHEET = 30
SAMPLE_SIZE = SHEETS_PER_WORKBOOK * ROWS_PER_SHEET
ANNOTATORS = ["annotator_1", "annotator_2", "annotator_3"]

HEADERS = [
    "question",
    "candidate_answer",
    "correct_answer",
    "wrong_answer",
    "label",
]


def clean_text(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def dataset_key_for_model(model: str) -> str:
    return "olmo" if model == "olmo32" else model


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def contradictory_result_path(results_dir: Path, model: str, run_id: str) -> Path:
    return results_dir / f"run_{run_id}_{model}_simple.json"


def build_annotation_rows(
    rq2_examples: Path,
    dataset_path: Path,
    contradictory_dir: Path,
    seed: int,
) -> list[dict[str, str]]:
    source_rows = read_csv_rows(rq2_examples)
    if SAMPLE_SIZE > len(source_rows):
        raise ValueError(f"sample size {SAMPLE_SIZE} exceeds available rows {len(source_rows)}")

    sampled_rows = random.Random(seed).sample(source_rows, SAMPLE_SIZE)
    dataset = load_json(dataset_path)
    result_cache: dict[tuple[str, str], dict] = {}
    annotation_rows = []

    for row in sampled_rows:
        model = row["model"]
        run_id = row["run_id"]
        qid = row["qid"]
        cache_key = (model, run_id)
        if cache_key not in result_cache:
            path = contradictory_result_path(contradictory_dir, model, run_id)
            if not path.exists():
                raise FileNotFoundError(path)
            result_cache[cache_key] = load_json(path)

        fact = dataset[dataset_key_for_model(model)][qid]
        result = result_cache[cache_key][qid][model]
        annotation_rows.append(
            {
                "question": clean_text(result.get("question") or fact.get("question")),
                "candidate_answer": clean_text(result.get("answer")),
                "correct_answer": clean_text(fact.get("obj")),
                "wrong_answer": clean_text(fact.get("false_obj")),
            }
        )

    return annotation_rows


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Arial", size=10)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin_gray)

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(HEADERS)).coordinate}"

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=ROWS_PER_SHEET + 1, max_col=len(HEADERS)):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 42,
        "B": 58,
        "C": 28,
        "D": 28,
        "E": 26,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 30
    for idx in range(2, ROWS_PER_SHEET + 2):
        ws.row_dimensions[idx].height = 78


def add_label_validation(ws) -> None:
    formula = '"' + ",".join(LABELS) + '"'
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showDropDown=False,
        errorTitle="Invalid label",
        error=f"Choose one of: {', '.join(LABELS)}",
        promptTitle="Human label",
        prompt=f"Choose one of: {', '.join(LABELS)}",
    )
    ws.add_data_validation(validation)
    validation.add(f"E2:E{ROWS_PER_SHEET + 1}")


def create_workbook(rows: list[dict[str, str]], annotator: str, output_dir: Path) -> Path:
    if len(rows) != SAMPLE_SIZE:
        raise ValueError(f"expected {SAMPLE_SIZE} sample rows, found {len(rows)}")

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_idx in range(SHEETS_PER_WORKBOOK):
        start = sheet_idx * ROWS_PER_SHEET
        end = start + ROWS_PER_SHEET
        ws = wb.create_sheet(f"Batch {sheet_idx + 1:02d}")
        ws.append(HEADERS)
        for row in rows[start:end]:
            ws.append(
                [
                    row.get("question", ""),
                    row.get("candidate_answer", ""),
                    row.get("correct_answer", ""),
                    row.get("wrong_answer", ""),
                    "",
                ]
            )
        add_label_validation(ws)
        style_sheet(ws)

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"rq2_human_evaluation_{annotator}.xlsx"
    wb.save(path)
    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create RQ2 annotation Excel workbooks.")
    parser.add_argument("--rq2-examples", type=Path, default=DEFAULT_RQ2_EXAMPLES)
    parser.add_argument("--dataset", type=Path, default=DEFAULT_DATASET)
    parser.add_argument("--contradictory-dir", type=Path, default=DEFAULT_CONTRADICTORY_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--annotators", nargs="+", default=ANNOTATORS)
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rows = build_annotation_rows(
        args.rq2_examples,
        args.dataset,
        args.contradictory_dir,
        args.seed,
    )
    outputs = [
        create_workbook(rows, annotator, args.output_dir)
        for annotator in args.annotators
    ]
    for output in outputs:
        print(f"[ok] wrote {output}")


if __name__ == "__main__":
    main()
